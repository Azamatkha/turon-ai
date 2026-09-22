"""
To'lov jadvalini .xlsx faylga yig'ish (mobil ilova uchun).

NEGA SERVERDA: vebda Excel brauzerda yig'iladi (frontend
`utils/paymentSchedule.ts`), lekin mobil ilova xuddi shu faylni o'zi yasamasin —
tayyor faylni `POST /v1/calculator/schedule/xlsx` dan oladi. Ko'rinish (ranglar,
ustunlar, "Jami", sug'urta qatorlari) veb faylga mos qilib yozilgan.

Yorliqlar 3 tilda — frontend `locales/*/chat.ts` dagi `sch*` kalitlari bilan
bir xil bo'lishi kerak.
"""

from io import BytesIO
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.calculator.schemas import ScheduleResult

ExcelLang = Literal["uz", "uz_cyrl", "ru"]

LABELS: dict[str, dict[str, str]] = {
    "uz": {
        "title": "To‘lov jadvali",
        "no": "№",
        "date": "Sana",
        "balance": "Kredit qoldig‘i",
        "principal": "Asosiy qarz",
        "interest": "Foizlar",
        "total": "Jami to‘lov",
        "days": "Kunlar",
        "total_row": "Jami",
        "insurance": "Sug‘urta bo‘yicha xarajatlar",
        "full_cost": "Kreditning to‘liq qiymati",
    },
    "uz_cyrl": {
        "title": "Тўлов жадвали",
        "no": "№",
        "date": "Сана",
        "balance": "Кредит қолдиғи",
        "principal": "Асосий қарз",
        "interest": "Фоизлар",
        "total": "Жами тўлов",
        "days": "Кунлар",
        "total_row": "Жами",
        "insurance": "Суғурта бўйича харажатлар",
        "full_cost": "Кредитнинг тўлиқ қиймати",
    },
    "ru": {
        "title": "График платежей",
        "no": "№",
        "date": "Дата",
        "balance": "Остаток кредита",
        "principal": "Основной долг",
        "interest": "Проценты",
        "total": "Всего к оплате",
        "days": "Дней",
        "total_row": "Итого",
        "insurance": "Расходы по страхованию",
        "full_cost": "Полная стоимость кредита",
    },
}

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Ranglar veb fayldagi (paymentSchedule.ts -> STYLES_XML) bilan bir xil
_BRAND = "1B4B7A"
_THIN = Side(style="thin", color="B7C4D4")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TOTAL_BORDER = Border(
    left=_THIN, right=_THIN, top=Side(style="medium", color=_BRAND), bottom=_THIN
)
_HDR_FILL = PatternFill("solid", fgColor=_BRAND)
_ZEBRA_FILL = PatternFill("solid", fgColor="F5F8FB")
_TOTAL_FILL = PatternFill("solid", fgColor="DCE6F1")
_NUM_FMT = "#,##0.00"
_CENTER = Alignment(horizontal="center", vertical="center")
_WIDTHS = [5.5, 12, 16.5, 18.5, 18.5, 18.5, 9.5]


def build_schedule_xlsx(res: ScheduleResult, lang: ExcelLang = "uz") -> bytes:
    """Jadval natijasidan .xlsx faylning baytlarini qaytaradi."""
    L = LABELS.get(lang, LABELS["uz"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Jadval"

    # Sarlavha (A:G birlashtirilgan)
    ws.merge_cells("A1:G1")
    ws["A1"] = L["title"]
    ws["A1"].font = Font(bold=True, size=15, color=_BRAND)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 8

    # Ustun sarlavhalari
    heads = [L[k] for k in ("no", "date", "balance", "principal", "interest", "total", "days")]
    for col, text in enumerate(heads, start=1):
        c = ws.cell(row=3, column=col, value=text)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HDR_FILL
        c.border = _BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 34

    # To'lovlar (juft qatorlar zebra fonda)
    r = 4
    for row in res.rows:
        zebra = row.k % 2 == 0
        values = [
            row.k,
            row.date.strftime("%d.%m.%Y"),
            round(row.balance, 2),
            round(row.principal, 2),
            round(row.interest, 2),
            round(row.total, 2),
            row.days,
        ]
        for col, value in enumerate(values, start=1):
            c = ws.cell(row=r, column=col, value=value)
            c.border = _BORDER
            if zebra:
                c.fill = _ZEBRA_FILL
            if col in (1, 2, 7):
                c.alignment = _CENTER
            else:
                c.number_format = _NUM_FMT
                c.alignment = Alignment(vertical="center")
            if col == 6:
                c.font = Font(bold=True)
        r += 1

    # "Jami" — yorliq A:C birlashtiriladi
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    totals = {
        1: L["total_row"],
        4: round(res.total_principal, 2),
        5: round(res.total_interest, 2),
        6: round(res.total_paid, 2),
    }
    for col in range(1, 8):
        c = ws.cell(row=r, column=col, value=totals.get(col))
        c.font = Font(bold=True, color=_BRAND)
        c.fill = _TOTAL_FILL
        c.border = _TOTAL_BORDER
        if col >= 4:
            c.number_format = _NUM_FMT
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[r].height = 24
    r += 2

    # Sug'urta va to'liq qiymat — yorliq A:E, qiymat F ustunida
    for label, value, bold in (
        (L["insurance"], res.insurance, False),
        (L["full_cost"], res.full_cost, True),
    ):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        font = Font(bold=True, color=_BRAND) if bold else Font()
        lbl = ws.cell(row=r, column=1, value=label)
        lbl.font = font
        lbl.alignment = Alignment(horizontal="left", vertical="center")
        num = ws.cell(row=r, column=6, value=round(value, 2))
        num.font = font
        num.number_format = _NUM_FMT
        ws.row_dimensions[r].height = 22
        r += 1

    for i, w in enumerate(_WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w

    # Sarlavha qatori muzlatiladi; chop etishda albom ko'rinishi, eniga sig'adi
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
